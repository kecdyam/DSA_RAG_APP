import os
import re
import random
import zipfile
import pymupdf
import docx   # comes from the python-docx package

from ocr_helper import ocr_page, ocr_image_bytes

root = "DATA"

OCR_FILES = {"HEAP_SORT.pdf", "Graph_Module6_Lecture5.pdf"}   # printed PDFs only
MIN_OCR_WORDS = 8   # OCR results with fewer real words are thrown away


def good_ocr(text):
    words = re.findall(r"[A-Za-z]{3,}", text)
    return len(words) >= MIN_OCR_WORDS


def good_image_ocr(text):
    code_like = re.findall(r"[A-Za-z][A-Za-z-]{3,}\(", text)   # e.g. PARTITION(
    words = re.findall(r"[A-Za-z]{4,}", text)
    return len(code_like) >= 1 or len(words) >= 12

def image_number(name):
    numbers = re.findall(r"\d+", os.path.basename(name))
    return int(numbers[0]) if numbers else 0


def clean_pdf_text(text):
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n", text)
    return text.strip()


def remove_license(text):
    text = text.lstrip("\ufeff")
    i = text.find("nbfs")
    if i != -1 and i < 600:
        start = text.rfind("/*", 0, i)
        end = text.find("*/", i)
        if start != -1 and end != -1:
            text = text[:start] + text[end + 2:]
    lines = [line for line in text.splitlines() if "nbfs://" not in line]
    return "\n".join(lines).strip()


def split_text(text, size=1000, overlap=150):
    chunks = []
    start = 0
    while start < len(text):
        end = start + size
        chunks.append(text[start:end])
        if end >= len(text):
            break
        start = end - overlap
    return chunks


all_chunks = []

for folder, subfolders, files in os.walk(root):
    for filename in files:
        if filename.startswith("~$"):      # temporary files Word leaves behind
            continue
        path = os.path.join(folder, filename)
        lower = filename.lower()

        if lower.endswith(".pdf"):
            doc = pymupdf.open(path)
            for page_number, page in enumerate(doc, start=1):
                raw = page.get_text()
                if filename in OCR_FILES and len(raw.strip()) < 50 and page.get_images():
                    ocr_text = ocr_page(page, filename, page_number)
                    if good_ocr(ocr_text):
                        raw = ocr_text
                text = clean_pdf_text(raw)
                for piece in split_text(text):
                    if len(piece) < 50:
                        continue
                    all_chunks.append({"text": piece, "source": filename,
                                       "page": page_number, "type": "pdf"})
            doc.close()

        elif lower.endswith(".java"):
            with open(path, encoding="utf-8-sig", errors="ignore") as f:
                text = remove_license(f.read())
            for piece in split_text(text):
                if len(piece) < 50:
                    continue
                all_chunks.append({"text": piece, "source": filename,
                                   "page": None, "type": "java"})

        elif lower.endswith(".docx"):
            d = docx.Document(path)
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for table in d.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            text = clean_pdf_text("\n".join(parts))
            for piece in split_text(text):
                if len(piece) < 50:
                    continue
                all_chunks.append({"text": piece, "source": filename,
                                   "page": None, "type": "docx"})

            # pictures inside the Word file (for example pseudocode screenshots)
            z = zipfile.ZipFile(path)
            image_names = sorted((n for n in z.namelist() if n.startswith("word/media/")),
                                 key=image_number)
            for n in image_names:
                ocr_text = ocr_image_bytes(z.read(n), f"{filename}_{os.path.basename(n)}")
                if good_image_ocr(ocr_text):
                    piece = f"Text read from a picture in {filename}:\n" + clean_pdf_text(ocr_text)
                    all_chunks.append({"text": piece, "source": filename,
                                       "page": None, "type": "docx"})

if __name__ == "__main__":
    print("Total chunks:", len(all_chunks))
    for chunk in random.sample(all_chunks, 3):
        print("-----", chunk["source"], "| page", chunk["page"])
        print(chunk["text"])